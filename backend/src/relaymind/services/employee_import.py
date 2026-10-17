"""Bounded XLSX preview and explicit synthetic directory snapshot import.

No HR data is sent to an LLM. No workbook formula is executed. This operator-only
service is not an authenticated public upload endpoint or an agent tool.
"""
import hashlib
import io
import json
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from relaymind.contracts.organization import CompanyDirectory
from relaymind.graph.access import AccessDenied, authorized_graph

# Company access is server-owned, never selected by a spreadsheet or a model.
COMPANY_SCOPES = {
    ("demo-a", "developer-a", "launch-01"): "northstar",
    ("demo-b", "developer-b", "launch-01"): "harbor",
    ("demo-c", "developer-c", "launch-01"): "cedar",
}
HEADERS = ("employee_id", "name", "team_id", "role", "manager_id", "status", "skills")
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024


def read_employee_xlsx(path: Path, company: CompanyDirectory) -> CompanyDirectory:
    """Replace a supplied company's roster with validated Employees-sheet rows.

    Caller supplies trusted company metadata; workbook cannot change company or
    graph. Skills are semicolon-separated; manager_id may be blank. Full roster
    only, not a partial update. Rejects unknown columns and unsupported formulas.
    """
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Expected .xlsx file no larger than 5 MiB")
    data = path.read_bytes()
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("Workbook expands beyond import limits")
            if any("vbaproject" in e.filename.lower() for e in entries):
                raise ValueError("Macro content is not supported")
    except BadZipFile:
        raise ValueError("Not a valid XLSX archive") from None
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Install the employee-import extra: pip install -e './backend[employee-import]'") from None
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if "Employees" not in workbook.sheetnames:
            raise ValueError("Workbook requires an Employees sheet")
        sheet = workbook["Employees"]
        if sheet.max_row and sheet.max_row > 1001:
            raise ValueError("At most 1000 employee rows are supported")
        if sheet.max_column and sheet.max_column != len(HEADERS):
            raise ValueError("Employees sheet must have exactly seven columns")
        rows = sheet.iter_rows()
        if tuple(cell.value for cell in next(rows, ())) != HEADERS:
            raise ValueError("Employees headers must match the documented seven-column template")
        employees = []
        for row_number, row in enumerate(rows, start=2):
            if row_number > 1001:
                raise ValueError("At most 1000 employee rows are supported")
            if any(c.data_type == "f" for c in row):
                raise ValueError(f"Formula in row {row_number}; use plain values")
            values = [c.value for c in row]
            if all(v is None for v in values):
                continue
            if len(values) != len(HEADERS) or any(v is not None and not isinstance(v, str) for v in values):
                raise ValueError(f"Row {row_number} requires text values, including employee IDs")
            record = dict(zip(HEADERS, (v.strip() if v else "" for v in values)))
            record["manager_id"] = record["manager_id"] or None
            record["skills"] = sorted(set(s.strip() for s in record["skills"].split(";") if s.strip()))
            employees.append(record)
        return CompanyDirectory.model_validate({**company.model_dump(), "employees": employees})
    finally:
        workbook.close()


def directory_identity(directory: CompanyDirectory):
    """Order-independent SHA-256 content identity. Never fuzzy-merge employees."""
    payload = directory.model_dump()
    payload["teams"] = sorted(payload["teams"], key=lambda t: t["id"])
    payload["employees"] = sorted(payload["employees"], key=lambda p: p["employee_id"])
    for person in payload["employees"]:
        person["skills"] = sorted(set(person["skills"]))
    content = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    digest = hashlib.sha256(content.encode()).hexdigest()
    return f"{directory.company_id}:{digest}", content


def save_directory(client, context, project_id, directory: CompanyDirectory, *, source_sha256: str | None = None):
    """Explicit operator apply: immutable snapshot, idempotent graph MERGE.

    Does not assign a current snapshot, alter task owners, or approve agent truth.
    Snapshot IDs must be selected explicitly on reads. Unique constraints and live
    concurrent-import tests are prerequisites for a multiworker deployment.
    """
    if COMPANY_SCOPES.get((context.tenant_id, context.user_id, project_id)) != directory.company_id:
        raise AccessDenied("Company access denied")
    graph = authorized_graph(client, context, project_id)
    snapshot_id, payload = directory_identity(directory)
    source_sha256 = source_sha256 or hashlib.sha256(payload.encode()).hexdigest()
    if len(source_sha256) != 64 or any(c not in "0123456789abcdef" for c in source_sha256):
        raise ValueError("Expected SHA-256 source fingerprint")
    teams = [{**t.model_dump(), "node_id": f"{snapshot_id}:team:{t.id}"} for t in directory.teams]
    people = [{**p.model_dump(), "node_id": f"{snapshot_id}:employee:{p.employee_id}",
               "team_node_id": f"{snapshot_id}:team:{p.team_id}"} for p in directory.employees]
    # One query: snapshot, all members, provenance rows and reporting edges.
    written = graph.query(
        "MERGE (c:Company {id:$company}) ON CREATE SET c.name=$name, c.synthetic=true "
        "MERGE (d:DirectorySnapshot {id:$snapshot}) "
        "ON CREATE SET d.payload_json=$payload, d.synthetic=true, d.project_id=$project "
        "MERGE (c)-[:HAS_DIRECTORY]->(d) "
        "MERGE (i:DirectoryImport {id:$import_id}) ON CREATE SET i.source_sha256=$source_sha "
        "MERGE (d)-[:SUPPORTED_BY]->(i) "
        "WITH d UNWIND $teams AS row "
        "MERGE (t:DirectoryTeam {id:row.node_id}) ON CREATE SET t.team_id=row.id, t.name=row.name "
        "MERGE (t)-[:IN_DIRECTORY]->(d) "
        "WITH DISTINCT d UNWIND $people AS row "
        "MATCH (t:DirectoryTeam {id:row.team_node_id}) "
        "MERGE (e:DirectoryEmployee {id:row.node_id}) "
        "ON CREATE SET e.employee_id=row.employee_id, e.name=row.name, e.role=row.role, "
        "e.status=row.status, e.skills=row.skills, e.manager_id=row.manager_id "
        "MERGE (e)-[:IN_DIRECTORY]->(d) MERGE (e)-[:IN_TEAM]->(t) "
        "WITH DISTINCT d UNWIND $people AS row "
        "MATCH (e:DirectoryEmployee {id:row.node_id}) "
        "OPTIONAL MATCH (m:DirectoryEmployee {id:$snapshot + ':employee:' + row.manager_id}) "
        "FOREACH (manager IN CASE WHEN m IS NULL THEN [] ELSE [m] END | MERGE (e)-[:REPORTS_TO]->(manager)) "
        "RETURN count(e)",
        params={"company": directory.company_id, "name": directory.company_name,
                "snapshot": snapshot_id, "payload": payload, "project": project_id,
                "source_sha": source_sha256, "import_id": snapshot_id + ':' + source_sha256,
                "teams": teams, "people": people}).result_set
    if written != [[len(directory.employees)]]:
        raise RuntimeError("Directory graph row count mismatch; do not activate this snapshot")
    stored = graph.ro_query("MATCH (d:DirectorySnapshot {id:$id}) RETURN d.payload_json",
                            params={"id": snapshot_id}).result_set
    if stored != [[payload]]:
        raise RuntimeError("Directory snapshot read-back failed")
    return snapshot_id


def load_directory(client, context, project_id, snapshot_id: str) -> CompanyDirectory:
    """Read an explicit snapshot only inside the caller's fixed company scope."""
    company = COMPANY_SCOPES.get((context.tenant_id, context.user_id, project_id))
    if company is None:
        raise AccessDenied("Company access denied")
    graph = authorized_graph(client, context, project_id)
    rows = graph.ro_query(
        "MATCH (:Company {id:$company})-[:HAS_DIRECTORY]->(d:DirectorySnapshot {id:$id, project_id:$project}) "
        "RETURN d.payload_json", params={"company": company, "id": snapshot_id, "project": project_id}).result_set
    if len(rows) != 1:
        raise AccessDenied("Directory unavailable or access denied")
    result = CompanyDirectory.model_validate_json(rows[0][0])
    if result.company_id != company or directory_identity(result)[0] != snapshot_id:
        raise RuntimeError("Directory snapshot integrity check failed")
    return result
